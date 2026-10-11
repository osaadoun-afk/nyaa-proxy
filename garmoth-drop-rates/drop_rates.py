"""Reverse-calculate per-trash (and per-kill) drop rates from Garmoth's collective grind data.

Input: JSON with {"collective": <api/grind-tracker/collective/all>, "spots": <api/grind-tracker/getGrindSpots>}.
Garmoth stores hourly_rate at 0% item drop rate and no loot buff, so ratios below are base rates.
Trash per kill varies by spot, so per-kill values are formulas driven by an editable trash/kill column.
"""
import csv
import json
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

LOW_HOURS = 50
Z = 1.96


def poisson_ci(n):
    """Byar's approximation of the 95% Poisson interval for an observed count n."""
    if n <= 0:
        return 0.0, 3.689  # upper bound for zero observed events
    lo = n * (1 - 1 / (9 * n) - Z / (3 * math.sqrt(n))) ** 3
    m = n + 1
    hi = m * (1 - 1 / (9 * m) + Z / (3 * math.sqrt(m))) ** 3
    return max(lo, 0.0), hi


def build(data):
    spots_meta = data["spots"]
    spots, drops = [], []
    for c in data["collective"]["data"]:
        meta = spots_meta.get(str(c["grindspot_id"]), {})
        hours = c["total_minutes"] / 60
        trash = next((d for d in c["drops"] if d["is_trash"]), None)
        trash_hr = trash["hourly_rate"] if trash else 0
        flag = "low sample" if hours < LOW_HOURS else ""
        spots.append({
            "spot_id": c["grindspot_id"], "spot": c["name"], "zone": meta.get("zone"),
            "ap": meta.get("ap"), "dp": meta.get("dp"), "sessions": c["total_sessions"],
            "hours": round(hours, 1), "weeks": c["total_weeks"],
            "trash_item": trash["name"] if trash else "", "base_trash_hr": trash_hr, "flag": flag,
        })
        for d in c["drops"]:
            if d["is_trash"]:
                continue
            rate = d["hourly_rate"]
            # Observed count lower bound: sessions ran with buffs, so real counts are >= base rate * hours.
            n_obs = rate * hours
            lo, hi = poisson_ci(n_obs)
            per_trash = rate / trash_hr if trash_hr else None
            scale = (per_trash / n_obs) if (per_trash and n_obs) else None
            drops.append({
                "spot_id": c["grindspot_id"], "spot": c["name"], "item": d["name"], "item_key": d["item_key"],
                "base_items_hr": rate, "base_trash_hr": trash_hr,
                "per_1000_trash": per_trash * 1000 if per_trash is not None else None,
                "one_in_n_trash": round(1 / per_trash) if per_trash else None,
                "ci_low_per_1000": lo * scale * 1000 if scale else 0.0,
                "ci_high_per_1000": (hi * scale * 1000) if scale else
                    (hi / hours / trash_hr * 1000 if trash_hr and hours else None),
                "est_drops_observed_min": round(n_obs, 1),
                "flag": "no drops recorded" if rate <= 0 else ("low sample" if n_obs < 30 or hours < LOW_HOURS else ""),
            })
    spots.sort(key=lambda s: s["spot"])
    drops.sort(key=lambda d: (d["spot"], -(d["per_1000_trash"] or 0)))
    return spots, drops


def write_xlsx(spots, drops, path, period):
    wb = Workbook()
    hdr = Font(bold=True)
    edit = PatternFill("solid", fgColor="FFF2CC")

    ws = wb.active
    ws.title = "Spots"
    cols = ["spot_id", "spot", "zone", "ap", "dp", "sessions", "hours", "weeks", "trash_item",
            "base_trash_hr", "trash_per_kill (EDIT)", "est_kills_hr", "flag"]
    ws.append(cols)
    for i, s in enumerate(spots, start=2):
        ws.append([s["spot_id"], s["spot"], s["zone"], s["ap"], s["dp"], s["sessions"], s["hours"], s["weeks"],
                   s["trash_item"], s["base_trash_hr"], None, f'=IF(K{i}="","",J{i}/K{i})', s["flag"]])
        ws[f"K{i}"].fill = edit
    spot_rows = len(spots) + 1

    wd = wb.create_sheet("Drops")
    dcols = ["spot_id", "spot", "item", "base_items_hr", "base_trash_hr", "per_1000_trash", "one_in_n_trash",
             "ci_low_per_1000", "ci_high_per_1000", "est_drops_observed_min",
             "trash_per_kill", "chance_per_kill_%", "one_in_n_kills", "flag"]
    wd.append(dcols)
    for i, d in enumerate(drops, start=2):
        tpk = f'IFERROR(INDEX(Spots!$K$2:$K${spot_rows},MATCH(A{i},Spots!$A$2:$A${spot_rows},0)),"")'
        wd.append([d["spot_id"], d["spot"], d["item"], d["base_items_hr"], d["base_trash_hr"],
                   d["per_1000_trash"], d["one_in_n_trash"], d["ci_low_per_1000"], d["ci_high_per_1000"],
                   d["est_drops_observed_min"],
                   f'=IF({tpk}="","",{tpk})',
                   f'=IF(OR(K{i}="",F{i}=""),"",F{i}/1000*K{i}*100)',
                   f'=IF(OR(L{i}="",L{i}=0),"",ROUND(100/L{i},0))',
                   d["flag"]])
    for c in "FHI":
        for cell in wd[c][1:]:
            cell.number_format = "0.0000"
    for cell in wd["L"][1:]:
        cell.number_format = "0.0000"

    wn = wb.create_sheet("Notes")
    for line in [
        f"Source: Garmoth collective grind data, {period}.",
        "All rates are base: 0% item drop rate, no loot scroll / Agris (Garmoth normalizes before storing).",
        "per_1000_trash = item hourly_rate / trash hourly_rate * 1000. Independent of clear speed and buffs.",
        "Trash per kill varies by spot, so per-kill chance needs Spots!K (yellow). Fill it in; Drops sheet updates.",
        "chance_per_kill_% = per_1000_trash / 1000 * trash_per_kill * 100. With Item Drop Rate buff, multiply by (1 + DR%).",
        "Items that drop in stacks: values are expected items, not drop events.",
        "CI: 95% Poisson interval on estimated observed drops (base rate * hours). Real samples are larger "
        "(sessions were buffed), so intervals are conservative.",
        f"Flags: 'low sample' = spot under {LOW_HOURS}h or fewer than 30 estimated drops; 'no drops recorded' = rate 0, "
        "ci_high gives an upper bound.",
    ]:
        wn.append([line])

    for sheet in (ws, wd):
        for cell in sheet[1]:
            cell.font = hdr
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for idx in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(idx)].width = 14
        sheet.column_dimensions["B"].width = 34
    wd.column_dimensions["C"].width = 40
    wn.column_dimensions["A"].width = 140
    wb.save(path)


def write_csv(rows, path):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def main():
    src, out = Path(sys.argv[1]), Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)
    data = json.loads(src.read_text())
    spots, drops = build(data)
    period = f'{data["collective"]["start_date"]} to {data["collective"]["end_date"]}'
    write_xlsx(spots, drops, out / "garmoth_drop_rates.xlsx", period)
    write_csv(drops, out / "drops.csv")
    write_csv(spots, out / "spots.csv")
    print(f"{len(spots)} spots, {len(drops)} drop rows -> {out}")


if __name__ == "__main__":
    main()
